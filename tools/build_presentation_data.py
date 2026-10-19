"""Дані для інтерактивної презентації (presentation/) і лонгріда (data/).

Читає xlsx «Бюджети_участі_2015_2025_статистичні_дані» (аркуші «Матриця 2015-2025»
і «Аналітика_частки») та JSON атласу (класи прозорості, голоси EGAP, картооснова),
пише presentation/data/deck.json, presentation/data/geo.json, а також
data/pb_data.json і data/analytics.json для лонгріда (той самий формат, що й раніше).

    python tools/build_presentation_data.py <xlsx> <atlas.json>

Числа, яких немає у xlsx (ціни 2021 р., таблиці 2–4, теми EGAP, правила конкурсів),
перенесено зі статті як константи нижче; скрипт звіряє з ними контрольні суми.
"""
import json
import statistics
import sys
from pathlib import Path

import openpyxl

ROOT = Path(__file__).resolve().parent.parent
OUT = ROOT / "presentation" / "data"
ART = ROOT / "data"
YEARS = list(range(2015, 2026))

# Макрорегіональні групи статті (табл. 2)
REGION = {
    "W": ["Луцьк", "Рівне", "Львів", "Тернопіль", "Івано-Франківськ", "Ужгород", "Чернівці", "Хмельницький"],
    "C": ["Київ", "Біла Церква", "Бровари", "Житомир", "Вінниця", "Чернігів", "Черкаси", "Суми", "Полтава",
          "Кременчук", "Кропивницький"],
    "SE": ["Одеса", "Миколаїв", "Херсон", "Запоріжжя", "Мелітополь", "Бердянськ", "Дніпро", "Кривий Ріг",
           "Кам’янське", "Нікополь", "Павлоград", "Харків", "Маріуполь", "Краматорськ", "Слов'янськ",
           "Сєвєродонецьк"],
}
OCCUPIED = {"Бердянськ", "Маріуполь", "Мелітополь", "Сєвєродонецьк"}
# Стан після 24.02.2022 (2023–2025), за текстом статті
POST_ACTIVE = {"Івано-Франківськ", "Луцьк", "Хмельницький", "Кам’янське"}
POST_INTAKE = {"Вінниця", "Тернопіль"}
POST_NOTE = {"Київ": "ГБ-7 перервано війною"}
# Суми з матриці, яких стаття не включає ні до документованих коштів, ні до пар частки
# (звірено з аркушем «Аналітика_частки» і рис. 2: 3,07 млрд грн, 113 пар)
EXCLUDE_AMOUNT = {("Черкаси", 2016), ("Черкаси", 2017), ("Черкаси", 2018), ("Маріуполь", 2020), ("Павлоград", 2021)}

TYPES = {
    "sv": ("стала практика, вагома вага",
           ["Івано-Франківськ", "Кривий Ріг", "Львів", "Рівне", "Тернопіль", "Ужгород", "Черкаси", "Чернівці", "Чернігів"]),
    "sn": ("стала практика, невелика вага",
           ["Вінниця", "Запоріжжя", "Кам’янське", "Київ", "Луцьк", "Хмельницький"]),
    "pv": ("переривчаста практика, вагома вага",
           ["Біла Церква", "Житомир", "Кременчук", "Миколаїв", "Одеса", "Полтава", "Суми", "Сєвєродонецьк", "Харків", "Херсон"]),
    "pn": ("переривчаста практика, невелика вага",
           ["Бердянськ", "Бровари", "Дніпро", "Краматорськ", "Кропивницький", "Мелітополь", "Нікополь", "Павлоград", "Слов'янськ"]),
    "na": ("без зіставних пар", ["Маріуполь"]),
}

# Рис. 2 статті: документовані кошти у цінах 2021 р., млн грн
REAL_2021 = {2015: 9, 2016: 234, 2017: 573, 2018: 666, 2019: 582, 2020: 671, 2021: 647,
             2022: 0, 2023: 24, 2024: 29, 2025: 43}

# Табл. 2 статті
GROUPS = [
    {"id": "W", "kind": "region", "label": "Захід", "cities": "8 / 8", "pairs": 30, "median": 0.36, "weighted": 0.38, "uah": 47.3, "apps": 0.29, "spend": 13.2},
    {"id": "C", "kind": "region", "label": "Центр і Північ", "cities": "11 / 11", "pairs": 34, "median": 0.30, "weighted": 0.29, "uah": 47.4, "apps": 0.31, "spend": 12.7},
    {"id": "SE", "kind": "region", "label": "Південь і Схід", "cities": "15 / 16", "pairs": 49, "median": 0.22, "weighted": 0.29, "uah": 20.5, "apps": 0.35, "spend": 11.0},
    {"id": "s1", "kind": "size", "label": "100–199 тис.", "cities": "9 / 9", "pairs": 25, "median": 0.20, "weighted": 0.20, "uah": 19.2, "apps": 0.49, "spend": None},
    {"id": "s2", "kind": "size", "label": "200–499 тис.", "cities": "18 / 19", "pairs": 61, "median": 0.32, "weighted": 0.33, "uah": 43.4, "apps": 0.32, "spend": None},
    {"id": "s3", "kind": "size", "label": "500–999 тис.", "cities": "4 / 4", "pairs": 16, "median": 0.30, "weighted": 0.30, "uah": 43.6, "apps": 0.27, "spend": None},
    {"id": "s4", "kind": "size", "label": "1 млн і більше", "cities": "3 / 3", "pairs": 11, "median": 0.32, "weighted": 0.30, "uah": 57.6, "apps": 0.19, "spend": None},
]

# Табл. 3 статті: done / fail / prog / unk
IMPLEMENTATION = [
    {"city": "Київ", "cohort": "ГБ-3 · подання 2018", "n": 341, "done": 316, "fail": 0, "prog": 0, "unk": 25, "asof": "фінальний звіт КМДА, 2020", "src": "агрегований звіт"},
    {"city": "Київ", "cohort": "ГБ-5 · подання 2020", "n": 351, "done": 340, "fail": 0, "prog": 0, "unk": 11, "asof": "01.01.2022", "src": "агрегований звіт"},
    {"city": "Біла Церква", "cohort": "переможці 2017–2020", "n": 40, "done": 28, "fail": 5, "prog": 7, "unk": 0, "asof": "29.11.2021", "src": "реєстр ради"},
    {"city": "Хмельницький", "cohort": "2021", "n": 23, "done": 16, "fail": 0, "prog": 0, "unk": 7, "asof": "03.07.2024", "src": "відкриті дані"},
    {"city": "Хмельницький", "cohort": "2025", "n": 17, "done": 16, "fail": 1, "prog": 0, "unk": 0, "asof": "26.01.2026", "src": "відкриті дані"},
    {"city": "Вінниця", "cohort": "БГІ-2021", "n": 23, "done": 14, "fail": 0, "prog": 1, "unk": 8, "asof": "01.10.2024", "src": "відкриті дані"},
    {"city": "Слов’янськ", "cohort": "2020 і 2021", "n": 15, "done": 15, "fail": 0, "prog": 0, "unk": 0, "asof": "01.01.2022", "src": "офіційний звіт"},
    {"city": "Кам’янське", "cohort": "2021, 2023, 2025", "n": 40, "done": 40, "fail": 0, "prog": 0, "unk": 0, "asof": "09.09.2026", "src": "відкриті дані"},
    {"city": "11 міст", "cohort": "«Громадський проєкт», 49 довоєнних когорт", "n": 1879, "done": 1446, "fail": 44, "prog": 389, "unk": 0, "asof": "23.07.2026", "src": "статуси платформи"},
    {"city": "Дніпро, Кривий Ріг", "cohort": "когорта 2022", "n": 154, "done": 0, "fail": 0, "prog": 154, "unk": 0, "asof": "23.07.2026", "src": "статуси платформи"},
]

# EGAP (2016–2019): частки тем серед поданих і переможців, %
THEMES = [
    {"id": "edu", "label": "Освіта", "submitted": 23, "won": 31},
    {"id": "public", "label": "Благоустрій", "submitted": 19, "won": 19},
    {"id": "sport", "label": "Спорт", "submitted": 14, "won": 15},
]

# Граничні кошториси проєктів, тис. грн (EGAP, 2020)
CAPS = [
    {"city": "Мелітополь", "year": "2019", "bands": [{"label": "малі", "lo": 0, "hi": 75}, {"label": "великі", "lo": 75, "hi": 350}]},
    {"city": "Суми", "year": "", "bands": [{"label": "локальні", "lo": 0, "hi": 750}, {"label": "загальноміські", "lo": 750, "hi": 1500}]},
    {"city": "Львів", "year": "", "bands": [{"label": "малі", "lo": 50, "hi": 600}, {"label": "великі", "lo": 600, "hi": 3000}]},
    {"city": "Київ", "year": "2017", "bands": [{"label": "малі", "lo": 100, "hi": 1000}, {"label": "великі", "lo": 1000, "hi": 3000}]},
]

# Медіана голосів на 1 тис. мешканців (EGAP), N міст
VOTE_MEDIANS = {2016: (47.8, 17), 2017: (66.1, None), 2018: (77.2, None), 2019: (89.8, 28)}


def num(v):
    if v is None or v == "":
        return None
    try:
        return float(v)
    except (TypeError, ValueError):
        return None


def read_matrix(ws):
    """Повертає {місто: {pop, first, src, years: {рік: {a, w, s, b}}}}."""
    header = [c.value for c in ws[3]]
    sub = [c.value for c in ws[4]]
    # Колонки кожного року: шукаємо за підписом року в рядку 3
    cols = {}
    year = None
    for i, (h, s) in enumerate(zip(header, sub)):
        if isinstance(h, (int, float)) and 2015 <= int(h) <= 2025:
            year = int(h)
        if year and s:
            key = {"Заявки / проєкти": "a", "Переможці": "w", "Сума БУ, грн": "s",
                   "Бюджет міста, грн": "b"}.get(str(s).strip())
            if key:
                cols.setdefault(year, {})[key] = i
    cities = {}
    for row in ws.iter_rows(min_row=5, values_only=True):
        name = row[0]
        if not name:
            continue
        years = {}
        for y in YEARS:
            c = cols.get(y, {})
            years[y] = {k: num(row[i]) for k, i in c.items()}
        cities[name.strip()] = {"pop": int(row[1]), "first": int(row[2]), "src": row[-1], "years": years}
    return cities


def read_share_analytics(ws):
    rows = list(ws.iter_rows(values_only=True))
    out = {}
    for r in rows:
        if isinstance(r[0], (int, float)) and 2018 <= int(r[0]) <= 2025 and r[1] is not None and r[7] is not None:
            out[int(r[0])] = {"n": int(r[1]), "mean": r[2] * 100, "median": r[3] * 100,
                              "weighted": r[4] * 100, "docmin": r[7] * 100}
    return out


def r2(x, nd=2):
    return None if x is None else round(x, nd)


def main(xlsx_path, atlas_path):
    wb = openpyxl.load_workbook(xlsx_path, data_only=True)
    matrix = read_matrix(wb.worksheets[0])
    share_by_year = read_share_analytics(wb.worksheets[1])
    atlas = json.loads(Path(atlas_path).read_text(encoding="utf-8"))
    acity = {c["name"]: c for c in atlas["cities"]}

    assert len(matrix) == 35, len(matrix)
    region_of = {c: g for g, cs in REGION.items() for c in cs}
    type_of = {c: k for k, (_, cs) in TYPES.items() for c in cs}
    assert set(region_of) == set(matrix), set(matrix) ^ set(region_of)
    assert set(type_of) == set(matrix), set(matrix) ^ set(type_of)

    cities = []
    clean = {}  # очищені роки міста (після вилучень) — для даних лонгріда
    for name, m in matrix.items():
        a = acity[name]
        ys = m["years"]
        clean[name] = ys
        # Заповненість бази «місто–рік» (табл. 1): заявки / переможці / сума, як у матриці
        mask = "".join("".join("1" if ys[y].get(k) else "0" for k in "aws") for y in YEARS)
        for y in YEARS:
            if (name, y) in EXCLUDE_AMOUNT:
                ys[y]["s"] = None
        # 2022: нових відборів не зафіксовано — повтори циклів 2021 р. і переможців вилучено,
        # залишаються лише заявки перерваних конкурсів
        v22, v21 = ys[2022], ys[2021]
        if v22.get("w") or (v22.get("a") and v22.get("a") == v21.get("a")):
            v22 = {}
        ys[2022] = {"a": v22.get("a"), "b": ys[2022].get("b")}
        pairs = [(ys[y]["s"], ys[y]["b"]) for y in range(2018, 2022) if ys[y].get("s") and ys[y].get("b")]
        apps = [ys[y]["a"] for y in range(2018, 2022) if ys[y].get("a")]
        share = sum(s for s, _ in pairs) / sum(b for _, b in pairs) * 100 if pairs else None
        uah = statistics.mean(s for s, _ in pairs) / m["pop"] if pairs else None
        app_k = statistics.mean(apps) / m["pop"] * 1000 if apps else None
        years_out = {}
        for y in YEARS:
            v = ys[y]
            years_out[y] = [v.get("a"), v.get("w"), r2(v["s"] / 1e6, 3) if v.get("s") else None]
        winners_years = sum(1 for y in YEARS if years_out[y][1])
        if name in OCCUPIED:
            post = "occupied"
        elif name in POST_ACTIVE:
            post = "active"
        elif name in POST_INTAKE:
            post = "intake"
        else:
            post = "none"
        size = ("s1" if m["pop"] < 200_000 else "s2" if m["pop"] < 500_000 else "s3" if m["pop"] < 1_000_000 else "s4")
        cities.append({
            "name": name, "oblast": a["oblast"], "lon": a["lon"], "lat": a["lat"],
            "pop": m["pop"], "first": m["first"], "region": region_of[name], "size": size,
            "type": type_of[name], "tr": a["transparency"], "trLabel": a["tlabel"],
            "occupied": name in OCCUPIED, "post": post, "postNote": POST_NOTE.get(name),
            "winYears": winners_years, "pairs": len(pairs),
            "share": r2(share, 3), "uah": r2(uah, 1), "appsK": r2(app_k, 2),
            "votes": {int(y): v["per1000"] for y, v in a["votes"].items()},
            "y": years_out, "mask": mask,
            "sh": {y: r2(ys[y]["s"] / ys[y]["b"] * 100, 3) for y in range(2018, 2022) if ys[y].get("s") and ys[y].get("b")},
        })
    cities.sort(key=lambda c: c["name"])

    annual = []
    for y in YEARS:
        cw = sum(1 for c in cities if c["y"][y][1])
        apps = sum(c["y"][y][0] or 0 for c in cities)
        win = sum(c["y"][y][1] or 0 for c in cities)
        amt = sum(c["y"][y][2] or 0 for c in cities)
        launched = sum(1 for c in cities if c["first"] <= y)
        sb = dict(share_by_year.get(y, {}))
        if y == 2022:  # рис. 1б: повтори циклів 2021 р. вилучено, нових відборів немає
            sb["docmin"] = 0.0
        annual.append({"year": y, "citiesWin": cw, "launched": launched, "newCities": sum(1 for c in cities if c["first"] == y),
                       "apps": int(apps), "winners": int(win), "nominal": round(amt, 1), "real": REAL_2021[y],
                       "docShare": r2(sb.get("docmin"), 3), "medShare": r2(sb.get("median"), 3),
                       "wShare": r2(sb.get("weighted"), 3), "nActive": sb.get("n")})

    tot_amt = sum(a["nominal"] for a in annual)
    tot_apps = sum(a["apps"] for a in annual)
    tot_win = sum(a["winners"] for a in annual)
    city_tot = sorted(((c["name"], sum(c["y"][y][2] or 0 for y in YEARS)) for c in cities), key=lambda t: -t[1])

    # Контрольні числа статті
    checks = {
        "Рис. 1а: міста з переможцями": ([a["citiesWin"] for a in annual], [2, 15, 23, 30, 27, 30, 30, 0, 3, 3, 4]),
        "Нові міста за роками": ([a["newCities"] for a in annual], [3, 17, 9, 4, 1, 1, 0, 0, 0, 0, 0]),
        "Заявки 2015–2025": (tot_apps, 26981),
        "Переможці 2015–2025": (tot_win, 6207),
        "Кошти 2015–2025, млрд": (round(tot_amt / 1000, 2), 3.07),
        "Частка Києва, %": (round(city_tot[0][1] / tot_amt * 100, 1), 28.6),
        "ТОП-5, %": (round(sum(v for _, v in city_tot[:5]) / tot_amt * 100, 1), 59.9),
        "Пари 2018–2021": (sum(c["pairs"] for c in cities), 113),
        "Позиції з трьома показниками": (sum(c["mask"][i:i + 3] == "111" for c in cities for i in range(0, 33, 3)), 161),
    }
    for k, (got, want) in checks.items():
        flag = "OK " if got == want else "DIFF"
        print(f"{flag} {k}: {got} (стаття: {want})")

    deck = {
        "meta": {"built": "2026-10-17", "source": Path(xlsx_path).name,
                 "frame": "35 міст: 34 міста понад 100 тис. осіб (01.01.2022) + Сєвєродонецьк"},
        "cities": cities,
        "annual": annual,
        "totals": {"apps": tot_apps, "winners": tot_win, "amountMln": round(tot_amt, 1),
                   "cityTotals": [{"name": n, "mln": round(v, 2)} for n, v in city_tot]},
        "groups": GROUPS,
        "types": {k: v[0] for k, v in TYPES.items()},
        "implementation": IMPLEMENTATION,
        "themes": THEMES,
        "caps": CAPS,
        "voteMedians": {y: {"median": m, "n": n} for y, (m, n) in VOTE_MEDIANS.items()},
    }
    OUT.mkdir(parents=True, exist_ok=True)
    (OUT / "deck.json").write_text(json.dumps(deck, ensure_ascii=False, separators=(",", ":")), encoding="utf-8")
    write_geo(atlas)
    print("Записано", OUT / "deck.json", "і geo.json")
    write_article_data(matrix, clean, cities)


def norm_name(n):
    return n.replace("'", "’")


def write_article_data(matrix, clean, cities):
    """data/pb_data.json і data/analytics.json лонгріда — з тієї самої бази, що й стаття.

    Зі старих файлів беруться лише незмінні атрибути: код області, координати, стани
    «призупинено» / «не проводився», основа суми (якщо сума збігається) та міжнародні кейси.
    """
    old = json.loads((ART / "pb_data.json").read_text(encoding="utf-8"))
    old_an = json.loads((ART / "analytics.json").read_text(encoding="utf-8"))
    oc = {norm_name(c["name"]): c for c in old["cities"]}
    occ = {c["name"]: c["occupied"] for c in cities}
    out_cities, obs = [], []
    for name, m in matrix.items():
        o, ys = oc[norm_name(name)], clean[name]
        years = {}
        for y in YEARS:
            v, ov = ys[y], o["years"].get(str(y), {})
            a, w, s, b = v.get("a"), v.get("w"), v.get("s"), v.get("b")
            if y < m["first"]:
                st = "pre"
            elif a or w or s:
                st = "data"
            elif ov.get("status") in ("susp", "notheld"):
                st = ov["status"]
            else:
                st = "nodata"
            same = s is not None and ov.get("s") is not None and abs(ov["s"] - s) < 1
            basis, prec = (ov.get("basis"), ov.get("precision")) if same else (None, None)
            share = s / b if (s and b and y >= 2018) else None
            years[str(y)] = {"a": a and int(a), "w": w and int(w), "s": s, "budget": b, "status": st,
                             "basis": basis, "precision": prec,
                             "perWinner": round(s / w) if s and w else None, "share": share}
            if share is not None:
                obs.append({"city": name, "year": y, "s": s, "budget": b, "share": share, "basis": basis, "precision": prec})
        out_cities.append({"name": name, "region": o["region"], "iso": o["iso"], "lon": o["lon"], "lat": o["lat"],
                           "population": m["pop"], "tot2022": occ[name], "firstYear": m["first"], "years": years})
    out_cities.sort(key=lambda c: c["name"])

    totals = {}
    for y in YEARS:
        rows = [c["years"][str(y)] for c in out_cities]
        totals[str(y)] = {
            "activeCities": sum(1 for r in rows if r["w"]),  # міста з визначеними переможцями (рис. 1а)
            "applications": int(sum(r["a"] or 0 for r in rows)),
            "winners": int(sum(r["w"] or 0 for r in rows)),
            "amount": round(sum(r["s"] or 0 for r in rows)),
            "citiesWithAmount": sum(1 for r in rows if r["s"]),
            "confirmedOnly": y >= 2022,
        }
    city_tot = sorted(({"name": c["name"], "amount": round(sum(r["s"] or 0 for r in c["years"].values()))} for c in out_cities),
                      key=lambda d: -d["amount"])
    pb = {
        "meta": {
            "title": "Бюджети участі великих міст України, 2015–2025",
            "generated": "2026-10-17",
            "frame": old["meta"]["frame"],
            "caveats": [
                "Порожня клітинка не дорівнює нулю: стани «немає даних», «до запуску», «призупинено», «не проводився» розрізняються явно.",
                "Суми мають різні основи (річний ліміт, сума бюджетів переможців, офіційний розподіл) і не усереднюються між містами.",
                "2022 р.: нових відборів не зафіксовано (повтори циклів 2021 р. вилучено); 2023–2025 — документований мінімум.",
            ],
        },
        "years": YEARS, "totals": totals, "cities": out_cities, "cityTotals": city_tot,
    }

    obs.sort(key=lambda d: (d["city"], d["year"]))

    def stats(rows):
        sh = [r["share"] for r in rows]
        return {"n": len(rows), "mean": statistics.mean(sh), "median": statistics.median(sh),
                "weighted": sum(r["s"] for r in rows) / sum(r["budget"] for r in rows)}

    by_year = [{"year": y, **stats([r for r in obs if r["year"] == y])} for y in YEARS if any(r["year"] == y for r in obs)]
    profiles = []
    for name in sorted({r["city"] for r in obs}):
        rows = [r for r in obs if r["city"] == name]
        sh = [r["share"] for r in rows]
        profiles.append({"city": name, "n": len(rows), "firstYear": rows[0]["year"], "lastYear": rows[-1]["year"],
                         "mean": statistics.mean(sh), "median": statistics.median(sh), "min": min(sh), "max": max(sh)})
    profiles.sort(key=lambda p: -p["mean"])
    pre = [r for r in obs if r["year"] <= 2021]
    post = [r for r in obs if r["year"] >= 2022]
    an = {
        "summary": {**stats(obs), "frame": "35 міст, 2018–2025"},
        "byYear": by_year,
        "observations": obs,
        "profiles": profiles,
        "periods": [{"label": "2018–2021", "n": len(pre), "mean": statistics.mean(r["share"] for r in pre)},
                    {"label": "2023–2025", "n": len(post), "mean": statistics.mean(r["share"] for r in post)}],
        "thresholds": old_an.get("thresholds", []),
        "international": old_an["international"],
        "barcelona": old_an["barcelona"],
    }
    s18 = stats(pre)
    print(f"Лонгрід: пар 2018–2021 = {s18['n']}, медіана {s18['median'] * 100:.3f}% (стаття 0,299), "
          f"зважена {s18['weighted'] * 100:.3f}% (стаття 0,306); усього спостережень {len(obs)}")
    (ART / "pb_data.json").write_text(json.dumps(pb, ensure_ascii=False, separators=(",", ":")), encoding="utf-8")
    (ART / "analytics.json").write_text(json.dumps(an, ensure_ascii=False, separators=(",", ":")), encoding="utf-8")
    print("Записано data/pb_data.json і data/analytics.json")


def rnd_line(line, nd=2):
    out = []
    for x, y in line:
        p = [round(x, nd), round(y, nd)]
        if not out or out[-1] != p:
            out.append(p)
    return out


def rnd_ring(ring, nd=2):
    """Округлене замкнене кільце; None, якщо після округлення воно виродилося."""
    out = rnd_line(ring, nd)
    if out and out[0] != out[-1]:
        out.append(out[0])
    return out if len(out) >= 4 else None


def rnd_poly(rings):
    out = [r for r in (rnd_ring(r) for r in rings) if r]
    return out if out and out[0] == rnd_ring(rings[0]) else None


def in_view(coords, box=(18.5, 42.5, 42.5, 54.2)):
    xs = [p[0] for p in coords]
    ys = [p[1] for p in coords]
    return not (max(xs) < box[0] or min(xs) > box[2] or max(ys) < box[1] or min(ys) > box[3])


def write_geo(atlas):
    ctx = atlas["context"]
    oblasts = []
    for f in atlas["geo"]["features"]:
        g = f["geometry"]
        polys = [g["coordinates"]] if g["type"] == "Polygon" else g["coordinates"]
        coords = [p for p in (rnd_poly(poly) for poly in polys) if p]
        oblasts.append({"name": f["properties"]["name"], "iso": f["properties"]["iso"], "c": coords})
    keep = lambda polys: [p for p in (rnd_poly(poly) for poly in polys if in_view(poly[0])) if p]
    geo = {
        "src": ctx["src"] + "; межі областей – geoBoundaries UKR ADM1, CC BY 4.0",
        "oblasts": oblasts,
        "land": keep(ctx["land"]),
        "lakes": keep(ctx["lakes"]),
        "borders": [rnd_line(l) for l in ctx["borders"] if in_view(l)],
        "uaBorder": [rnd_line(l) for l in ctx["ua_border"]],
        "uaCoast": [rnd_line(l) for l in ctx["ua_coast"]],
        "rivers": [{"n": r["n"], "w": r["w"], "g": rnd_line(r["g"])} for r in ctx["rivers"] if in_view(r["g"])],
        "labels": ctx["labels"],
    }
    (OUT / "geo.json").write_text(json.dumps(geo, ensure_ascii=False, separators=(",", ":")), encoding="utf-8")


if __name__ == "__main__":
    main(sys.argv[1], sys.argv[2])
